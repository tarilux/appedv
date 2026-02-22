from io import BytesIO
from django.db.models import Sum
from django.utils import timezone
from django.http import HttpResponse
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from pptx import Presentation

from .models import Project


def export_projects_excel(projects):
    wb = Workbook()
    ws = wb.active
    ws.title = "PROtrack Projects"
    ws.append(["Référence", "Intitulé", "État", "Budget prévisionnel", "Montant contrat", "Avancement", "Archivé"])

    for project in projects:
        ws.append([
            project.reference,
            project.intitule,
            project.etat,
            float(project.budget_previsionnel),
            "" if not project.montant_contrat_affichage else float(project.montant_contrat),
            project.taux_avancement,
            "Oui" if project.archive else "Non",
        ])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="protrack_projects.xlsx"'
    return response


def export_project_pdf(project):
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.setTitle(f"Fiche_{project.reference}")

    lines = [
        "PROtrack - Fiche Projet",
        f"Référence: {project.reference}",
        f"Intitulé: {project.intitule}",
        f"État: {project.etat}",
        f"Budget prévisionnel: {project.budget_previsionnel} €",
        f"Montant contrat: {project.montant_contrat_affichage or '-'}",
        f"Avancement: {project.taux_avancement}%",
        f"Date début: {project.date_debut}",
        f"Date fin prévue: {project.date_fin_prevue}",
        f"Description: {project.description or '-'}",
    ]

    y = 800
    for line in lines:
        pdf.drawString(50, y, line)
        y -= 25
    pdf.save()

    buffer.seek(0)
    response = HttpResponse(buffer, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="fiche_{project.reference}.pdf"'
    return response


def export_dashboard_ppt():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "PROtrack - Synthèse Dashboard"

    active = Project.objects.filter(archive=False)
    total = active.count()
    en_execution = active.filter(etat=Project.Etat.EXECUTION).count()
    clotures = active.filter(etat=Project.Etat.CLOTURE).count()
    budget_total = active.aggregate(total=Sum("budget_previsionnel"))["total"] or 0

    body = slide.shapes.placeholders[1].text_frame
    body.text = f"Date: {timezone.localdate()}"
    body.add_paragraph().text = f"Projets actifs: {total}"
    body.add_paragraph().text = f"En exécution: {en_execution}"
    body.add_paragraph().text = f"Clôturés: {clotures}"
    body.add_paragraph().text = f"Budget total: {budget_total} €"

    buffer = BytesIO()
    prs.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    response["Content-Disposition"] = 'attachment; filename="protrack_dashboard.pptx"'
    return response
