import json
from io import BytesIO

from django.contrib import messages
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .forms import ProjectForm, SOUS_ETAPES_PAR_ETAPE, SITUATION_OPTIONS
from .models import HistoriqueModification, Project


TRACKED_FIELDS = [
    'id_projet', 'type', 'exercice', 'etape_en_cours', 'sous_etape', 'libelle', 'chef_projet',
    'montant_estime', 'montant_engage', 'num_contrat', 'montant_contrat', 'fournisseur',
    'date_signature_os', 'delai_projet', 'date_fin_contrat', 'date_fin_actualisee', 'autonomie',
    'total_attachement', 'dernier_mois_attache', 'avancement_travaux', 'travaux_precedents',
    'travaux_en_cours', 'travaux_prochains', 'situation_projet', 'statut'
]


def serialize_value(value):
    return '' if value is None else str(value)


def capture_values(project):
    return {field: serialize_value(getattr(project, field)) for field in TRACKED_FIELDS}


def log_differences(project, before, after):
    for field in TRACKED_FIELDS:
        old = before.get(field, '')
        new = after.get(field, '')
        if old != new:
            HistoriqueModification.objects.create(
                projet=project,
                champ_modifie=field,
                ancienne_valeur=old,
                nouvelle_valeur=new,
            )


def get_active_filtered_queryset(request):
    queryset = Project.objects.filter(statut='actif').order_by('-updated_at')
    filter_etape = request.GET.get('etape', '').strip()
    filter_chef = request.GET.get('chef', '').strip()
    if filter_etape:
        queryset = queryset.filter(etape_en_cours=filter_etape)
    if filter_chef:
        queryset = queryset.filter(chef_projet=filter_chef)
    return queryset


def compute_dashboard_metrics(queryset):
    totals = queryset.aggregate(total_engage=Sum('montant_engage'), total_receptionne=Sum('total_attachement'))
    etape_counts, chef_counts, engagement_vs_paiement, reception_percent_values = {}, {}, [], []

    for p in queryset:
        etape_counts[p.etape_en_cours] = etape_counts.get(p.etape_en_cours, 0) + 1
        chef_counts[p.chef_projet] = chef_counts.get(p.chef_projet, 0) + 1
        engagement_vs_paiement.append({'label': p.id_projet, 'engage': float(p.montant_engage), 'receptionne': float(p.total_attachement)})
        if p.pourcentage_reception is not None:
            reception_percent_values.append(float(p.pourcentage_reception))

    avg_reception = round(sum(reception_percent_values) / len(reception_percent_values), 2) if reception_percent_values else 0
    return {
        'total_engage': totals['total_engage'] or 0,
        'total_receptionne': totals['total_receptionne'] or 0,
        'avg_reception': avg_reception,
        'etape_counts': etape_counts,
        'chef_counts': chef_counts,
        'engagement_vs_paiement': engagement_vs_paiement,
    }


def dashboard(request):
    show_archived = request.GET.get('archived') == '1'
    active_queryset = get_active_filtered_queryset(request)
    projects = Project.objects.all().order_by('-updated_at') if show_archived else active_queryset
    metrics = compute_dashboard_metrics(active_queryset)

    return render(request, 'projects/dashboard.html', {
        'projects': projects,
        'show_archived': show_archived,
        'archived_count': Project.objects.filter(statut='archive').count(),
        'active_count': active_queryset.count(),
        'total_engage': metrics['total_engage'],
        'total_receptionne': metrics['total_receptionne'],
        'avg_reception': metrics['avg_reception'],
        'chart_etape_labels': json.dumps(list(metrics['etape_counts'].keys()), ensure_ascii=False),
        'chart_etape_values': json.dumps(list(metrics['etape_counts'].values())),
        'chart_chef_labels': json.dumps(list(metrics['chef_counts'].keys()), ensure_ascii=False),
        'chart_chef_values': json.dumps(list(metrics['chef_counts'].values())),
        'chart_evsp_labels': json.dumps([x['label'] for x in metrics['engagement_vs_paiement']], ensure_ascii=False),
        'chart_evsp_engage': json.dumps([x['engage'] for x in metrics['engagement_vs_paiement']]),
        'chart_evsp_receptionne': json.dumps([x['receptionne'] for x in metrics['engagement_vs_paiement']]),
        'etape_options': Project.ETAPE_CHOICES,
        'chef_options': Project.objects.filter(statut='actif').values_list('chef_projet', flat=True).distinct().order_by('chef_projet'),
        'selected_etape': request.GET.get('etape', '').strip(),
        'selected_chef': request.GET.get('chef', '').strip(),
        'querystring': request.GET.urlencode(),
    })


def projects_page(request):
    queryset = Project.objects.filter(statut='actif').order_by('-updated_at')
    data = []
    for p in queryset:
        pct = None if p.pourcentage_reception is None else float(p.pourcentage_reception)
        data.append({
            'id': p.id,
            'id_projet': p.id_projet,
            'type': p.type,
            'exercice': p.exercice,
            'etape': p.etape_en_cours,
            'sous_etape': p.sous_etape,
            'libelle': p.libelle,
            'chef': p.chef_projet,
            'montant_contrat': float(p.montant_contrat),
            'total_attachement': float(p.total_attachement),
            'pourcentage_reception': pct,
            'avancement': float(p.avancement_travaux),
            'autonomie': p.autonomie,
            'situation': p.situation_projet or '',
        })
    return render(request, 'projects/projects_page.html', {
        'projects_json': json.dumps(data, ensure_ascii=False),
        'type_options': [c[0] for c in Project.TYPE_CHOICES],
        'exercice_options': [c[0] for c in Project.EXERCICE_CHOICES],
        'etape_options': [c[0] for c in Project.ETAPE_CHOICES],
        'sous_etape_options': sorted(set(Project.objects.filter(statut='actif').values_list('sous_etape', flat=True))),
        'chef_options': sorted(set(Project.objects.filter(statut='actif').values_list('chef_projet', flat=True))),
        'autonomie_options': sorted(set(Project.objects.filter(statut='actif').values_list('autonomie', flat=True))),
        'situation_options': sorted(set(Project.objects.filter(statut='actif').exclude(situation_projet__isnull=True).exclude(situation_projet='').values_list('situation_projet', flat=True))),
    })


def project_detail(request, pk):
    project = get_object_or_404(Project, pk=pk)
    return render(request, 'projects/project_detail.html', {'project': project})


def project_create(request):
    if request.method == 'POST':
        form = ProjectForm(request.POST)
        if form.is_valid():
            project = form.save()
            log_differences(project, {}, capture_values(project))
            messages.success(request, 'Projet créé avec succès.')
            return redirect('projects_page')
    else:
        form = ProjectForm()
    return render(request, 'projects/project_form.html', {
        'form': form,
        'title': 'Nouveau projet',
        'project': None,
        'sous_etapes_map': json.dumps(SOUS_ETAPES_PAR_ETAPE, ensure_ascii=False),
        'situation_options': json.dumps(SITUATION_OPTIONS, ensure_ascii=False),
    })


def project_edit(request, pk):
    project = get_object_or_404(Project, pk=pk)
    if request.method == 'POST':
        before = capture_values(project)
        form = ProjectForm(request.POST, instance=project)
        if form.is_valid():
            project = form.save()
            log_differences(project, before, capture_values(project))
            messages.success(request, 'Projet mis à jour.')
            return redirect('projects_page')
    else:
        form = ProjectForm(instance=project)
    return render(request, 'projects/project_form.html', {
        'form': form,
        'title': 'Modifier projet',
        'project': project,
        'sous_etapes_map': json.dumps(SOUS_ETAPES_PAR_ETAPE, ensure_ascii=False),
        'situation_options': json.dumps(SITUATION_OPTIONS, ensure_ascii=False),
    })


@require_POST
def project_archive(request, pk):
    project = get_object_or_404(Project, pk=pk)
    before = capture_values(project)
    project.statut = 'archive'
    project.save(update_fields=['statut', 'updated_at'])
    log_differences(project, before, capture_values(project))
    messages.warning(request, 'Projet archivé.')
    return redirect('projects_page')


def archive_list(request):
    projects = Project.objects.filter(statut='archive').order_by('-updated_at')
    return render(request, 'projects/archive.html', {'projects': projects})


@require_POST
def project_restore(request, pk):
    project = get_object_or_404(Project, pk=pk)
    before = capture_values(project)
    project.statut = 'actif'
    project.save(update_fields=['statut', 'updated_at'])
    log_differences(project, before, capture_values(project))
    messages.success(request, 'Projet restauré.')
    return redirect('archive_list')


def history_list(request):
    logs = HistoriqueModification.objects.select_related('projet').all()[:500]
    return render(request, 'projects/history.html', {'logs': logs})


def project_rows(queryset):
    for p in queryset.order_by('id_projet'):
        yield [
            p.id_projet, p.type, p.exercice, p.etape_en_cours, p.sous_etape, p.libelle, p.chef_projet,
            float(p.montant_estime), float(p.montant_engage), p.num_contrat, float(p.montant_contrat),
            p.fournisseur, p.date_signature_os.isoformat() if p.date_signature_os else '', p.delai_projet,
            p.date_fin_contrat, p.date_fin_actualisee, p.autonomie, float(p.total_attachement),
            '' if p.pourcentage_reception is None else float(p.pourcentage_reception),
            p.dernier_mois_attache, float(p.avancement_travaux), p.statut,
        ]


def export_xlsx(request):
    queryset = get_active_filtered_queryset(request)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Projets'
    ws.append([
        'ID Projet', 'Type', 'Exercice', 'Étape', 'Sous-étape', 'Libellé', 'Chef Projet', 'Montant Estimé',
        'Montant Engagé', 'N° Contrat', 'Montant Contrat', 'Fournisseur', 'Date Signature OS', 'Délai',
        'Fin Contrat', 'Fin Actualisée', 'Autonomie', 'Total Attachement', '% Réception',
        'Dernier Mois Attaché', 'Avancement Travaux', 'Statut',
    ])
    for row in project_rows(queryset):
        ws.append(row)
    out = BytesIO()
    wb.save(out)
    response = HttpResponse(out.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=projets_filtres.xlsx'
    return response


def export_pdf(request):
    queryset = get_active_filtered_queryset(request)
    metrics = compute_dashboard_metrics(queryset)

    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4)
    styles = getSampleStyleSheet()

    kpi_data = [
        ['KPI', 'Valeur'],
        ['Nombre projets actifs', str(queryset.count())],
        ['Total engagé', str(metrics['total_engage'])],
        ['Total réceptionné', str(metrics['total_receptionne'])],
        ['% moyen réception', f"{metrics['avg_reception']} %"],
    ]
    kpi_table = Table(kpi_data, colWidths=[180, 250], repeatRows=1)
    kpi_table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f2937')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke), ('GRID', (0, 0), (-1, -1), 0.5, colors.grey)]))

    etape_data = [['Étape', 'Volume']] + [[k, v] for k, v in metrics['etape_counts'].items()]
    chef_data = [['Chef projet', 'Volume']] + [[k, v] for k, v in metrics['chef_counts'].items()]
    summary_data = [['ID', 'Libellé', 'Engagé', 'Réceptionné']] + [
        [p.id_projet, p.libelle, p.montant_engage, p.total_attachement] for p in queryset[:12]
    ]

    story = [
        Paragraph('Rapport synthèse - Suivi Projets', styles['Heading2']),
        Spacer(1, 8),
        Paragraph('KPI', styles['Heading3']),
        kpi_table,
        Spacer(1, 10),
        Paragraph('Graphiques (données de répartition)', styles['Heading3']),
        Table(etape_data, repeatRows=1),
        Spacer(1, 6),
        Table(chef_data, repeatRows=1),
        Spacer(1, 10),
        Paragraph('Tableau synthèse', styles['Heading3']),
        Table(summary_data, repeatRows=1),
    ]
    doc.build(story)

    response = HttpResponse(out.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename=rapport_synthese.pdf'
    return response


def export_pptx(request):
    queryset = get_active_filtered_queryset(request)
    metrics = compute_dashboard_metrics(queryset)

    prs = Presentation()

    # Slide dashboard KPI
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.title.text = 'Dashboard - Projets filtrés'
    t1 = s1.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(2.5)).text_frame
    t1.text = f"Actifs: {queryset.count()} | Engagé: {metrics['total_engage']} | Réceptionné: {metrics['total_receptionne']} | % moyen: {metrics['avg_reception']}%"
    t1.paragraphs[0].font.size = Pt(18)

    # Slide graphiques (données)
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = 'Graphiques (répartitions)'
    t2 = s2.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(5)).text_frame
    t2.text = 'Répartition par étape:'
    for k, v in metrics['etape_counts'].items():
        p = t2.add_paragraph()
        p.text = f"- {k}: {v}"
    p = t2.add_paragraph()
    p.text = 'Répartition par chef de projet:'
    for k, v in metrics['chef_counts'].items():
        p = t2.add_paragraph()
        p.text = f"- {k}: {v}"

    # Slide tableau synthèse
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = 'Tableau synthèse'
    rows = min(8, queryset.count()) + 1
    cols = 4
    table_shape = s3.shapes.add_table(rows, cols, Inches(0.4), Inches(1.3), Inches(9.2), Inches(4.8))
    table = table_shape.table
    headers = ['ID', 'Libellé', 'Engagé', 'Réceptionné']
    for c, h in enumerate(headers):
        table.cell(0, c).text = h
    for i, p in enumerate(queryset[: rows - 1], start=1):
        table.cell(i, 0).text = p.id_projet
        table.cell(i, 1).text = p.libelle
        table.cell(i, 2).text = str(p.montant_engage)
        table.cell(i, 3).text = str(p.total_attachement)

    out = BytesIO()
    prs.save(out)
    response = HttpResponse(out.getvalue(), content_type='application/vnd.openxmlformats-officedocument.presentationml.presentation')
    response['Content-Disposition'] = 'attachment; filename=dashboard_projets.pptx'
    return response
